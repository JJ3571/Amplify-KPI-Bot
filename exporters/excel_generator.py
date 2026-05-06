"""
Excel Generator - Create KPI Calculator workbook
Generates Excel workbook with Export-KPI and Export-QA tabs
matching the structure in examples/[v2.4] KPI Calculator.xlsx
"""

import pandas as pd
import logging
from pathlib import Path
from datetime import datetime
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows


class ExcelGenerator:
    """Generate KPI Calculator Excel workbook from aggregated data"""
    
    def __init__(self, output_dir='data/exports'):
        """
        Initialize Excel generator
        
        Args:
            output_dir (str): Directory to save generated Excel files
        """
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        
        # Styling constants
        self.HEADER_FILL = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        self.HEADER_FONT = Font(bold=True, color='FFFFFF', size=11)
        self.BORDER_THIN = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def generate_workbook(self, export_kpi: pd.DataFrame, export_qa: pd.DataFrame,
                         week_start: datetime, week_end: datetime,
                         missing_data: list = None) -> Optional[str]:
        """
        Generate complete KPI Calculator workbook
        
        Args:
            export_kpi (pd.DataFrame): KPI data for Export-KPI tab
            export_qa (pd.DataFrame): QA data for Export-QA tab
            week_start (datetime): Week start date
            week_end (datetime): Week end date
            missing_data (list): List of data sources that are missing (for notes)
        
        Returns:
            str: Path to generated Excel file or None if failed
        """
        try:
            # Generate filename
            start_str = week_start.strftime('%Y-%m-%d')
            end_str = week_end.strftime('%Y-%m-%d')
            filename = f"KPI_Calculator_{start_str}_{end_str}.xlsx"
            filepath = self.output_dir / filename
            
            logging.info(f"Generating Excel workbook: {filename}")
            
            # Create workbook
            wb = Workbook()
            
            # Remove default sheet
            if 'Sheet' in wb.sheetnames:
                wb.remove(wb['Sheet'])
            
            # Create Export-KPI tab
            if export_kpi is not None:
                self._create_export_kpi_sheet(wb, export_kpi)
            
            # Create Export-QA tab
            if export_qa is not None:
                self._create_export_qa_sheet(wb, export_qa)
            
            # Create Notes sheet if there's missing data
            if missing_data:
                self._create_notes_sheet(wb, week_start, week_end, missing_data)
            
            # Save workbook
            wb.save(filepath)
            logging.info(f"✅ Excel workbook saved: {filepath}")
            
            return str(filepath)
            
        except Exception as e:
            logging.error(f"Failed to generate Excel workbook: {e}")
            return None
    
    def _create_export_kpi_sheet(self, wb: Workbook, df: pd.DataFrame):
        """Create Export-KPI sheet with formatting"""
        ws = wb.create_sheet('Export - KPI')
        
        # Write headers
        headers = list(df.columns)
        ws.append(headers)
        
        # Style header row
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = self.BORDER_THIN
        
        # Write data rows
        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=False), 2):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = value
                cell.border = self.BORDER_THIN
                
                # Format percentage columns
                col_name = headers[c_idx - 1]
                if '%' in col_name or 'Utilization' in col_name:
                    if isinstance(value, (int, float)) and value < 1:
                        cell.number_format = '0.00%'
                    else:
                        cell.number_format = '0.00'
                
                # Format numeric columns
                elif 'Hours' in col_name or 'min' in col_name or 'sec' in col_name:
                    cell.number_format = '0.00'
                
                # Format integer columns
                elif 'Cases' in col_name or 'SLA' in col_name or 'Responses' in col_name:
                    cell.number_format = '0'
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Freeze top row
        ws.freeze_panes = 'A2'
        
        logging.info(f"  Created Export-KPI sheet: {len(df)} rows, {len(headers)} columns")
    
    def _create_export_qa_sheet(self, wb: Workbook, df: pd.DataFrame):
        """Create Export-QA sheet with formatting"""
        ws = wb.create_sheet('Export - QA')
        
        # Write headers
        headers = list(df.columns)
        ws.append(headers)
        
        # Style header row
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = self.BORDER_THIN
        
        # Write data rows
        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=False), 2):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = value
                cell.border = self.BORDER_THIN
                
                # Format percentage columns
                col_name = headers[c_idx - 1]
                if '%' in col_name:
                    if isinstance(value, (int, float)) and value < 1:
                        cell.number_format = '0.00%'
                    else:
                        cell.number_format = '0.00'
                
                # Format integer columns
                elif 'Cases' in col_name or 'Audited' in col_name:
                    cell.number_format = '0'
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Freeze top row
        ws.freeze_panes = 'A2'
        
        logging.info(f"  Created Export-QA sheet: {len(df)} rows, {len(headers)} columns")
    
    def _create_notes_sheet(self, wb: Workbook, week_start: datetime, week_end: datetime,
                           missing_data: list):
        """Create Notes sheet with information about missing data"""
        ws = wb.create_sheet('Notes', 0)  # Insert as first sheet
        
        # Title
        ws['A1'] = 'KPI Calculator - Data Generation Notes'
        ws['A1'].font = Font(bold=True, size=14)
        
        # Week info
        week_str = f"{week_start.strftime('%m/%d/%Y')} - {week_end.strftime('%m/%d/%Y')}"
        ws['A3'] = f'Week: {week_str}'
        ws['A3'].font = Font(bold=True)
        
        # Generation date
        ws['A4'] = f'Generated: {datetime.now().strftime("%m/%d/%Y %I:%M %p")}'
        
        # Data availability
        ws['A6'] = 'Data Source Availability:'
        ws['A6'].font = Font(bold=True)
        
        row = 7
        sources = {
            'Tableau QA': 'QA scores and audit data',
            'Tableau KPI': 'CSAT, response times, case metrics',
            'Google Sheets': 'SLA cases by day',
            'Salesforce': 'Cases closed, cases transferred'
        }
        
        for source, description in sources.items():
            if any(src in source for src in missing_data):
                ws[f'A{row}'] = f'❌ {source}'
                ws[f'A{row}'].font = Font(color='FF0000')
                ws[f'B{row}'] = f'Missing - {description}'
            else:
                ws[f'A{row}'] = f'✅ {source}'
                ws[f'A{row}'].font = Font(color='00B050')
                ws[f'B{row}'] = description
            row += 1
        
        # Missing data instructions
        if missing_data:
            ws[f'A{row + 1}'] = 'Manual Data Entry Required:'
            ws[f'A{row + 1}'].font = Font(bold=True, color='FF0000')
            
            row += 2
            if 'Google Sheets' in ' '.join(missing_data) or 'SLA' in ' '.join(missing_data):
                ws[f'A{row}'] = '• SLA Cases columns (Sunday - Saturday) need to be manually populated'
                row += 1
            
            if 'Salesforce' in ' '.join(missing_data):
                ws[f'A{row}'] = '• Cases Closed and Cases Transferred need to be manually populated'
                row += 1
            
            row += 1
            ws[f'A{row}'] = 'Copy/paste this data into your live KPI Calculator Google Sheet'
            ws[f'A{row}'].font = Font(italic=True)
        
        # Auto-adjust column widths
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 60
        
        logging.info("  Created Notes sheet")
    
    def generate_csv_exports(self, export_kpi: pd.DataFrame, export_qa: pd.DataFrame,
                            week_start: datetime, week_end: datetime) -> tuple:
        """
        Generate separate CSV files for Export-KPI and Export-QA
        
        Args:
            export_kpi (pd.DataFrame): KPI data
            export_qa (pd.DataFrame): QA data
            week_start (datetime): Week start date
            week_end (datetime): Week end date
        
        Returns:
            tuple: (kpi_path, qa_path) - paths to generated CSV files
        """
        try:
            start_str = week_start.strftime('%Y-%m-%d')
            end_str = week_end.strftime('%Y-%m-%d')
            
            kpi_path = None
            qa_path = None
            
            # Save KPI CSV
            if export_kpi is not None:
                kpi_filename = f"Export_KPI_{start_str}_{end_str}.csv"
                kpi_path = self.output_dir / kpi_filename
                export_kpi.to_csv(kpi_path, index=False)
                logging.info(f"✅ Export-KPI CSV saved: {kpi_path}")
            
            # Save QA CSV
            if export_qa is not None:
                qa_filename = f"Export_QA_{start_str}_{end_str}.csv"
                qa_path = self.output_dir / qa_filename
                export_qa.to_csv(qa_path, index=False)
                logging.info(f"✅ Export-QA CSV saved: {qa_path}")
            
            return (str(kpi_path) if kpi_path else None, 
                   str(qa_path) if qa_path else None)
            
        except Exception as e:
            logging.error(f"Failed to generate CSV exports: {e}")
            return (None, None)


# Standalone usage
if __name__ == "__main__":
    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s - %(levelname)s - %(message)s'
    )
    
    print("="*70)
    print("Excel Generator - Testing")
    print("="*70)
    print()
    
    # Create sample data for testing
    from datetime import datetime, timedelta
    
    # Calculate last week
    today = datetime.now()
    days_since_sunday = (today.weekday() + 1) % 7
    last_sunday = today - timedelta(days=days_since_sunday + 7)
    last_saturday = last_sunday + timedelta(days=6)
    
    week_str = f"{last_sunday.strftime('%m/%d/%Y')} - {last_saturday.strftime('%m/%d/%Y')}"
    
    # Sample Export-QA data
    export_qa = pd.DataFrame({
        'Week': [week_str] * 3,
        'Agent Name': ['John Doe', 'Jane Smith', 'Bob Johnson'],
        'Team': ['Orange', 'Purple', 'Orange'],
        'BCF %': [0.03, 0.05, 0.02],
        'QA Score %': [0.95, 0.92, 0.97],
        'Cases Audited': [10, 12, 8]
    })
    
    # Sample Export-KPI data
    export_kpi = pd.DataFrame({
        'Week': [week_str] * 3,
        'Agent Name': ['John Doe', 'Jane Smith', 'Bob Johnson'],
        'Time Utilization %': [0.88, 0.92, 0.85],
        'Total Hours': [35.2, 36.8, 34.0],
        'Working Hours': [40.0, 40.0, 40.0],
        'CSAT Responses': [15, 18, 12],
        'CSAT Positive': [13, 16, 11],
        'Chat Response (min)': [4.5, 3.8, 5.2],
        'Phone Response (sec)': [45.0, 38.0, 52.0],
        'Email Response (min)': [420.0, 385.0, 450.0],
        'Cases Closed': [42, 48, 38],
        'Cases Transferred': [5, 3, 7],
        'SLA Sunday': [0, 1, 0],
        'SLA Monday': [2, 0, 1],
        'SLA Tuesday': [1, 1, 0],
        'SLA Wednesday': [0, 2, 1],
        'SLA Thursday': [1, 0, 2],
        'SLA Friday': [0, 1, 0],
        'SLA Saturday': [0, 0, 1]
    })
    
    # Initialize generator
    generator = ExcelGenerator()
    
    # Test 1: Generate Excel workbook
    print("📊 Generating Excel workbook...")
    excel_path = generator.generate_workbook(
        export_kpi, export_qa, last_sunday, last_saturday,
        missing_data=['Salesforce']  # Simulate missing Salesforce data
    )
    
    if excel_path:
        print(f"✅ Excel workbook created: {excel_path}")
    else:
        print("❌ Failed to create Excel workbook")
    
    # Test 2: Generate CSV exports
    print("\n📊 Generating CSV exports...")
    kpi_csv, qa_csv = generator.generate_csv_exports(
        export_kpi, export_qa, last_sunday, last_saturday
    )
    
    if kpi_csv and qa_csv:
        print(f"✅ CSV exports created:")
        print(f"   - {kpi_csv}")
        print(f"   - {qa_csv}")
    else:
        print("❌ Failed to create CSV exports")
    
    print("\n" + "="*70)
